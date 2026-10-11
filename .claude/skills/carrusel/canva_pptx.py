"""Genera un .pptx 1080x1350 con texto editable para importar en Canva.

Uso: python3 .claude/skills/carrusel/canva_pptx.py carruseles/NN-slug/canva.json
Sale carruseles/NN-slug/canva.pptx

canva.json: {"firma": "@usuario", "fondo": "#2B1D1F", "texto": "#FFFFFF", "acento": "#6E1A2E",
             "slides": [[["cuerpo", "Texto con **negrita**"], ["frase", "Frase ==resaltada=="], ["puente", "Sigue →"]], ...]}
Tipos de bloque: cuerpo (sans), frase (serif bold grande), titulo (serif bold más grande), puente (serif cursiva),
palabra (serif bold en caja clara, para la palabra del CTA).
"""
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525  # EMU por píxel
W, H, MARGEN = 1080, 1350, 104
SERIF, SANS = "Noto Serif Display", "Quicksand"
ESTILOS = {  # tamaño en px, fuente, negrita, cursiva, espacio antes en px
    "cuerpo": (44, SANS, False, False, 36),
    "frase": (76, SERIF, True, False, 40),
    "titulo": (104, SERIF, True, False, 44),
    "puente": (52, SERIF, False, True, 56),
    "palabra": (80, SERIF, True, False, 20),
}


def rgb(hexa):
    return RGBColor.from_string(hexa.lstrip("#").upper())


def runs(texto):
    """Parte '**negrita**' y '==resaltado==' en tramos (texto, negrita, resaltado)."""
    for parte in re.split(r"(\*\*.+?\*\*|==.+?==)", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            yield parte[2:-2], True, False
        elif parte.startswith("=="):
            yield parte[2:-2], True, True
        else:
            yield parte, False, False


def resaltar(run, color):
    rpr = run._r.get_or_add_rPr()
    hl = etree.SubElement(rpr, qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr"), val=color.lstrip("#").upper())


def caja(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def main(ruta):
    datos = json.loads(Path(ruta).read_text(encoding="utf-8"))
    texto, acento, fondo = datos.get("texto", "#FFFFFF"), datos["acento"], datos["fondo"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    total = len(datos["slides"])
    for i, bloques in enumerate(datos["slides"], 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(fondo)

        tb, tf = caja(s, MARGEN, 120, W - 2 * MARGEN, H - 300)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tb.name = "Texto"
        for j, (tipo, contenido) in enumerate(bloques):
            tam, fuente, negrita, cursiva, antes = ESTILOS[tipo]
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.2 if tipo in ("frase", "titulo") else 1.3
            if j:
                p.space_before = Pt(antes * 0.75)
            for frag, b, hl in runs(contenido):
                r = p.add_run()
                r.text = frag
                f = r.font
                f.name, f.size = fuente, Pt(tam * 0.75)
                f.bold, f.italic = negrita or b, cursiva or (hl and tipo in ("frase", "titulo"))
                f.color.rgb = rgb(acento if tipo == "palabra" else texto)
                if hl:
                    resaltar(r, acento)
                if tipo == "palabra":
                    resaltar(r, texto)

        for x, txt, alin in ((MARGEN, datos["firma"].upper(), PP_ALIGN.LEFT),
                             (W - MARGEN - 200, f"{i}/{total}", PP_ALIGN.RIGHT)):
            _, pf = caja(s, x, H - 92, 600 if alin == PP_ALIGN.LEFT else 200, 36)
            p = pf.paragraphs[0]
            p.alignment = alin
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = SANS, Pt(22 * 0.75), True
            r.font.color.rgb = rgb(texto)
    out = Path(ruta).with_name("canva.pptx")
    prs.save(out)
    print(f"PPTX exportado: {out} ({total} slides)")


if __name__ == "__main__":
    main(sys.argv[1])
